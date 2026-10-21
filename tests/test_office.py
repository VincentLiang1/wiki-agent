"""office-dump:攝入時核對 Office 原件用(原文字形、顏色、公式)。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Color, Font, PatternFill

from wiki_feeder.office import dump


def test_xlsx_shows_original_text_colors_formulas_and_blank_template_rows(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "送禮名單"
    ws["A1"] = "已聯絡"
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["B1"] = "統計名片"
    ws["B1"].fill = PatternFill("solid", fgColor=Color(theme=6, tint=0.6))
    ws["A2"] = "游定森 董事長"                         # doc2md 會繁化成「遊」
    ws["B2"] = "=SUM(C2:C3)"
    for col in "ABC":
        ws[f"{col}5"].font = Font(color="0000FF")      # 範本預設格式、沒有內容
    p = tmp_path / "名單.xlsx"
    wb.save(p)

    out = dump(p)

    assert "A2: '游定森 董事長'" in out
    assert "A1: '已聯絡'|底色 #FFFF00" in out
    assert "佈景主題色 6" in out                        # 只認 RGB 會漏掉這種
    assert "公式 =SUM(C2:C3)" in out
    assert "另有 3 個空白儲存格只有預設格式" in out


def test_docx_and_pptx_text(tmp_path):
    import docx
    import pptx

    d = docx.Document()
    d.add_paragraph("游定森 董事長致詞")
    d.save(tmp_path / "a.docx")
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "公司簡介"
    prs.save(tmp_path / "b.pptx")

    assert "游定森 董事長致詞" in dump(tmp_path / "a.docx")
    assert "公司簡介" in dump(tmp_path / "b.pptx")
    assert "不支援" in dump(tmp_path / "c.xls")


def test_dump_command_cannot_be_shadowed_from_the_repo_cwd(tmp_path):
    """攝入時 claude 的 cwd 是 FWIKI、也能寫檔:`python -m wiki_feeder` 會先找 cwd,
    它寫一個 wiki_feeder.py 再跑白名單裡這行,就等於能執行任意程式(push、刪原件)。"""
    import subprocess

    from wiki_feeder.config import INGEST_TOOLS_CMD

    (tmp_path / "wiki_feeder.py").write_text(
        "open('pwned', 'w').write('x')\n", encoding="utf-8")
    (tmp_path / "a.txt").write_text("a", encoding="utf-8")
    cmd = INGEST_TOOLS_CMD   # Config.default() 與設定檔都用它;Windows 沒設定檔時 default() 會報錯
    subprocess.run([*cmd, "office-dump", str(tmp_path / "a.txt")], cwd=tmp_path, capture_output=True, timeout=60)
    assert not (tmp_path / "pwned").exists()


def test_dump_labels_follow_the_interface_language_but_cell_text_does_not(tmp_path):
    from wiki_feeder import i18n
    wb = Workbook()
    ws = wb.active
    ws.title = "名單"
    ws["A1"] = "已聯絡"
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    p = tmp_path / "名單.xlsx"
    wb.save(p)
    i18n.set_lang("en")
    out = dump(p)
    assert out.startswith("# 名單.xlsx (original text, not converted to Traditional Chinese)")
    assert "## Sheet: 名單 (1 rows × 1 columns)" in out
    assert "A1: '已聯絡'|fill #FFFF00" in out
    assert dump(tmp_path / "x.rtf").startswith(".rtf is not supported")
