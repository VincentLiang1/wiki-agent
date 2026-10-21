"""Office 原件的只讀傾印:給攝入時的 claude 核對 doc2md 轉不出來的東西。

doc2md 轉出的 md 會繁化人名(「遊定森」其實是「游定森」這類要回原件核對),
也看不到儲存格的底色、字色與公式。這台 Mac 沒有裝 Office,改用 openpyxl／python-docx／
python-pptx 直接讀檔——只讀、不寫回,攝入時的白名單只開放這一個指令(不開放任意 python)。
"""
from __future__ import annotations

from pathlib import Path

from .i18n import L


def _color(c) -> str | None:
    """openpyxl 的顏色物件 → 可讀字串;預設色(黑字、系統預設索引色)回 None。

    ⚠️ 三種都要認:只認 RGB 會漏掉佈景主題色與索引色——實測一份名單 RGB 只抓到 1 格,
    其餘標註全是佈景主題色。
    """
    if c is None:
        return None
    kind = getattr(c, "type", None)
    if kind == "rgb" and isinstance(c.rgb, str):
        rgb = c.rgb[-6:].upper()
        return None if rgb == "000000" or c.rgb.upper() == "00000000" else f"#{rgb}"
    if kind == "theme" and isinstance(c.theme, int):
        tint = L("、明暗 {tint}", tint=f"{c.tint:+.2f}") if getattr(c, "tint", 0) else ""
        return None if c.theme == 1 and not tint else L("佈景主題色 {n}{tint}", n=c.theme, tint=tint)
    if kind == "indexed" and isinstance(c.indexed, int):
        return None if c.indexed in (8, 64) else L("索引色 {n}", n=c.indexed)   # 8 黑、64 系統預設
    return None


def dump_xlsx(path: Path) -> list[str]:
    import warnings

    import openpyxl

    warnings.filterwarnings("ignore", module="openpyxl")  # 「Data Validation extension」之類的提示不影響讀值

    values = openpyxl.load_workbook(path, data_only=True)    # 公式算出來的值
    formulas = openpyxl.load_workbook(path, data_only=False)  # 公式本身
    out = []
    for ws in values.worksheets:
        fws = formulas[ws.title]
        out.append(L("## 工作表:{title}({rows} 列 × {cols} 欄)", title=ws.title, rows=ws.max_row,
                      cols=ws.max_column))
        blank_styled = []
        for row in ws.iter_rows():
            for cell in row:
                f = fws[cell.coordinate].value
                fill = _color(cell.fill.fgColor) if cell.fill and cell.fill.fill_type == "solid" else None
                font = _color(cell.font.color) if cell.font else None
                if cell.value is None and not (isinstance(f, str) and f.startswith("=")):
                    if fill or font:
                        blank_styled.append(cell.coordinate)
                    continue
                parts = [f"{cell.coordinate}: {cell.value!r}"]
                if isinstance(f, str) and f.startswith("="):
                    parts.append(L("公式 {formula}", formula=f))
                if fill:
                    parts.append(L("底色 {color}", color=fill))
                if font:
                    parts.append(L("字色 {color}", color=font))
                if cell.comment:
                    parts.append(L("註解 {text}", text=repr(cell.comment.text.strip())))
                out.append("  " + "|".join(parts))
        if blank_styled:
            # 範本預先設好格式、還沒填內容的列:不是標註(doc2md 會把它們也算進「以顏色標註的儲存格」)
            out.append(L("  (另有 {n} 個空白儲存格只有預設格式、沒有內容:{cells}{more})", n=len(blank_styled),
                         cells=L("、").join(blank_styled[:6]), more=' …' if len(blank_styled) > 6 else ''))
    return out


def dump_docx(path: Path) -> list[str]:
    import docx

    d = docx.Document(str(path))
    out = [L("## 段落")]
    for i, p in enumerate(d.paragraphs, 1):
        if p.text.strip():
            marks = {r.font.highlight_color for r in p.runs if r.font.highlight_color}
            extra = L("|螢光標記 {marks}", marks=sorted(str(m) for m in marks)) if marks else ""
            out.append(f"  {i}[{p.style.name}]: {p.text}{extra}")
    for t, table in enumerate(d.tables, 1):
        out.append(L("## 表格 {n}", n=t))
        for r, row in enumerate(table.rows, 1):
            out.append(f"  {r}: " + " | ".join(c.text.strip() for c in row.cells))
    return out


def dump_pptx(path: Path) -> list[str]:
    import pptx

    prs = pptx.Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(L("## 第 {n} 頁", n=i))
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(f"  [{shape.name}] {shape.text_frame.text}")
            if getattr(shape, "has_table", False) and shape.has_table:
                for r, row in enumerate(shape.table.rows, 1):
                    out.append(L("  表格 {n}: ", n=r) + " | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(L("  備忘稿:{text}", text=slide.notes_slide.notes_text_frame.text))
    return out


DUMPERS = {".xlsx": dump_xlsx, ".xlsm": dump_xlsx, ".docx": dump_docx, ".docm": dump_docx,
           ".pptx": dump_pptx, ".pptm": dump_pptx}


def dump(path: Path) -> str:
    fn = DUMPERS.get(path.suffix.lower())
    if fn is None:
        return L("不支援 {suffix}(只支援 {supported};舊版 .doc/.xls/.ppt 請以 doc2md 的轉檔為準)", suffix=path.suffix,
                 supported=L("、").join(sorted(DUMPERS)))
    return "\n".join([L("# {name}(原件原文,未經繁化)", name=path.name), *fn(path)])
